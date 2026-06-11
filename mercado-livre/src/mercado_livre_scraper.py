# Mercado Livre - Web Scraping Automation

### Imports
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from webdriver_manager.chrome import ChromeDriverManager
from selenium.webdriver.support.wait import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

from time import sleep

import pandas as pd

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo


### Navigating to the page
service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service)

driver.get("https://www.mercadolivre.com.br/ofertas")
driver.maximize_window()
wait = WebDriverWait(driver, timeout=5)
wait.until(EC.visibility_of_element_located((By.XPATH, "//div[contains(@class, 'carousel_item')]")))

### Functions
# Function 1: Save DataFrames to Excel

def Saving_DataFrames(path, dfs):
    with pd.ExcelWriter(path) as writer:
        for category_name, df in dfs.items():
            df.to_excel(writer, sheet_name=category_name, index=False)


# function 2: DataFrame Customization

def Excel_Customizing(path):
    wb = load_workbook(path)

    for i, ws in enumerate(wb.worksheets):
    
        table_ref = ws.dimensions

        table_name = f"Table_{i}"

        table = Table(displayName=table_name, ref=table_ref)

        style = TableStyleInfo(
            name="TableStyleLight15", 
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        table.tableStyleInfo = style

        ws.add_table(table)

        col_idx = dfs[category_name].columns.get_loc("Price") + 1
        col_letter = get_column_letter(col_idx)

        rule = ColorScaleRule(
            start_type="min",
            start_color="FF2962A7", # Light Blue
            mid_type="percentile",
            mid_value=50,
            mid_color="FFFFFF", # White
            end_type="max",
            end_color ="FFFD3939" # Light Red
        )

        ws.conditional_formatting.add(f"{col_letter}2:{col_letter}{ws.max_row}", rule)

    wb.save(path)


### Web Scraping
carrosseis = driver.find_elements(By.XPATH, "//div[contains(@class, 'carousel_item')]")

dfs = {}

for i in range(len(carrosseis)):
    carrossel = carrosseis[i]

    while not carrossel.is_displayed():
        next_button = driver.find_element(By.XPATH, "//button[contains(@class, 'andes-carousel-snapped__control--next')]")
        next_button.click()
        sleep(2)
        carrosseis = driver.find_elements(By.CLASS_NAME, 'carousel_item')
        carrossel = carrosseis[i]
       
    category_name = carrossel.find_element(By.XPATH, ".//p[@class='title']").text

    carrossel.click()
    sleep(1)

    carrosseis = driver.find_elements(By.CLASS_NAME, 'carousel_item')
    carrossel = carrosseis[i]
    
    d = []
    s = []
    p = []

    itens = driver.find_elements(By.XPATH, "//div[@class='poly-card__content']")
    for item in itens:

        xpath_description = ".//a[contains(@class,'poly-component__title')]"
        description = item.find_element(By.XPATH, xpath_description).text
        d.append(description)

        xpath_seller = ".//span[@class='poly-component__seller']"
        seller = ("Por Mercado Livre" 
                    if len(item.find_elements(By.XPATH, xpath_seller)) == 0 
                    else item.find_element(By.XPATH, xpath_seller).text)
        s.append(seller)
        
        xpath_reais = ".//span[contains(@class,'andes-money-amount andes-money-amount--cents-superscript')]//span[contains(@class, 'andes-money-amount__fraction')]"
        reais = item.find_element(By.XPATH, xpath_reais).text
        reais = reais.replace(".","")

        xpath_cents = ".//span[contains(@class,'andes-money-amount andes-money-amount--cents-superscript')]//span[contains(@class, 'andes-money-amount__cents')]"
        cents = ("00" 
                    if len(item.find_elements(By.XPATH, xpath_cents)) == 0
                    else item.find_element(By.XPATH, xpath_cents).text)
        
        price = float(f"{reais}.{cents}")
        p.append(price)

        print(category_name, " || ", description," | ", seller," | ",price)

    dfs[category_name] = pd.DataFrame({
    'Product': d,
    'Seller': s,
    'Price': p
})
    dfs[category_name] = dfs[category_name].sort_values(by='Price', ascending=False).reset_index(drop=True)


### Saving DataFrames
path = Path(__file__).parent.parent / 'output' / 'OfertasMeli.xlsx'
Saving_DataFrames(path, dfs)


### Customizing DataFrames
Excel_Customizing(path)