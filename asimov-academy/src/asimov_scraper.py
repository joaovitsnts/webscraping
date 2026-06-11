# Asimov Academy - Web Scraping Automation


### Imports
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from webdriver_manager.chrome import ChromeDriverManager
from selenium.webdriver.support.wait import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.action_chains import ActionChains
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Alignment
from time import sleep
from datetime import datetime
from pathlib import Path
import pandas as pd


### Navigating to the page
service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service)

driver.get("https://hub.asimov.academy")
driver.maximize_window()


### Login Credentials Input
wait = WebDriverWait(driver, timeout=10)
email_input = wait.until(EC.visibility_of_element_located((By.ID, "email")))
email_input.send_keys("your_email")
password_input = driver.find_element(By.ID, "password")
password_input.send_keys("your_password")
entrar_button = driver.find_element(By.XPATH, "//button[text()='Entrar']")
entrar_button.click()


### Accessing My Certifications Section
avatar_menu = wait.until(EC.visibility_of_element_located((By.ID, "userMenuTrigger")))
avatar_menu.click()
certificados_area = driver.find_element(By.XPATH, "//a[contains(@title,'Ir para certificados') and contains(.,'Certificados')]")
certificados_area.click()


### Getting Certificates
certificados_finished = driver.find_element(By.XPATH, "//div[contains(@class,'flex-col gap-6')][.//h2[contains(.,'Cursos')]]")
list_names = certificados_finished.find_elements(By.XPATH, ".//*[@class='text-xs text-neutral-950']")
course_name = [name.text for name in list_names]
list_dates = certificados_finished.find_elements(By.XPATH, ".//span[contains(@class,'text-neutral-600') and contains(text(), 'Obtido em:')]")
finish_date = [date.text.replace('Obtido em: ','') for date in list_dates]
finish_date = [datetime.strptime(date, "%d/%m/%y") for date in finish_date]


### Viewing All Courses
course_button = driver.find_element(By.XPATH, "//a[contains(@href, 'cursos')]")
course_button.click()
page_course = wait.until(EC.visibility_of_element_located((By.XPATH, "//div[@id='contentWrapper']")))
titles = page_course.find_elements(By.XPATH, "//h2")
titles = [title.text for title in titles]
status = ["Completed" if title in course_name else "Not Completed" for title in titles]


### Functions
def save_dataframes(all_courses, finished_courses, finish_dates, save_path):
    df_finished = pd.DataFrame({
        "Courses": finished_courses,
        "Finished On": finish_dates
    })

    all_course_status = status = ["Completed" if title in finished_courses else "Not Completed" for title in all_courses]
    df_all_courses = pd.DataFrame({
        'Courses': all_courses,
        'Status': all_course_status
    })

    with pd.ExcelWriter(path / 'Status_asimov_academy.xlsx') as writer:
        df_all_courses.to_excel(writer, sheet_name='All Courses', index=False)
        df_finished.to_excel(writer, sheet_name='Finished Courses', index=False)
        
    print('DataFrame successfully saved')

def customizing(path, sheet_amount=1):
    
    wb = load_workbook(path / "status_asimov_academy.xlsx")

    header_fill = PatternFill(
        start_color="778899",
        end_color="778899",
        fill_type="solid"
    )

    header_font = Font(bold=True, color="FFFFFF")
    alignment = Alignment(horizontal="center", vertical="center")
    no_border = Border()
    font = Font(size=10)

    for i in range(sheet_amount):

        sheet = wb.worksheets[i]

        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = alignment
                cell.border = no_border
                cell.font = font

        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = alignment

    wb.save(path / "Status_asimov_academy.xlsx")


### Saving Data Frame
path = Path(__file__).parent.parent / 'output'

if __name__ == "__main__":
    save_dataframes(titles, course_name, finish_date, path)
    customizing(path, 2)